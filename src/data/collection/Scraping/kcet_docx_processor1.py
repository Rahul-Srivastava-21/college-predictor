import os
import re
import csv
import logging
from docx import Document
from pathlib import Path

# ------------------ Config ------------------
INPUT_DIR = r"D:\Courses\Global Academy of Technology\kcet-college-pred\src\data\collection\Scraping\kcet_docs"
OUTPUT_FILE = "kcet_cutoffs_all_streamed.csv"
DEBUG_DIR = "debug_tables"

logging.basicConfig(level=logging.INFO, format="%(message)s")

# ------------------ Optional progress bar (tqdm fallback) ------------------
try:
    from tqdm import tqdm
except Exception:
    # fallback tqdm
    def tqdm(iterable, **kwargs):
        total = kwargs.get("total", None)
        desc = kwargs.get("desc", "")
        for i, item in enumerate(iterable, 1):
            if desc:
                print(f"{desc} [{i}/{total if total is not None else '?'}]")
            yield item

# ------------------ Helper Functions ------------------

def extract_code_num(code: str):
    m = re.match(r"E(\d{3})", code)
    return int(m.group(1)) if m else None

def clean_cell_text(text: str):
    return " ".join(text.split()) if text else ""

def dump_debug(table, code, reason):
    os.makedirs(DEBUG_DIR, exist_ok=True)
    safe_code = code.replace("/", "_").replace("\\", "_")
    fname = f"{safe_code}_{reason}.csv"
    path = os.path.join(DEBUG_DIR, fname)
    with open(path, "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        for row in table.rows:
            writer.writerow([clean_cell_text(c.text) for c in row.cells])
    logging.warning(f"⚠️ Dumped suspicious table for {code} → {path}")
    return path

def iter_blocks(doc, filepath):
    """
    Iterate over paragraphs and tables in a DOCX.
    If we encounter a <w:t> with None text, we log a warning and
    abort processing this file (Option 2).
    """
    for block in doc.element.body:
        if block.tag.endswith("p"):
            texts = []
            for t in block.iter():
                if t.tag.endswith("t"):
                    if t.text is None:
                        logging.warning(f"⚠️ Skipping file {filepath} due to empty <w:t> element")
                        return  # stop processing this file
                    texts.append(t.text)
            text = "".join(texts).strip()
            yield ("p", text)
        elif block.tag.endswith("tbl"):
            for t in doc.tables:
                if t._element == block:
                    yield ("tbl", t)
                    break

# ------------------ Main Processing ------------------

def process_docx(filepath, year, round_num):
    logging.info(f"\n📂 Processing {filepath}...")
    doc = Document(filepath)

    rows_written = 0
    tables_processed = 0
    debug_dumps = 0
    placeholders = 0
    colleges_found = 0
    warnings = []

    seen_rows = set()
    last_code_num = None
    current_college_code = None
    current_college_name = None

    with open(OUTPUT_FILE, "a", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)

        # use iter_blocks with filepath for logging context
        for kind, content in iter_blocks(doc, filepath) or []:
            if kind == "p":
                match = re.match(r"(E\d{3})\s+(.+)", content)
                if match:
                    current_college_code, current_college_name = match.groups()
                    colleges_found += 1
                    current_code_num = extract_code_num(current_college_code)

                    if last_code_num is not None and current_code_num is not None and current_code_num != last_code_num + 1:
                        missing = list(range(last_code_num + 1, current_code_num))
                        for miss in missing:
                            fake_code = f"E{miss:03d}"
                            msg = f"❌ Missing expected college {fake_code} (detected gap {last_code_num} -> {current_code_num})"
                            logging.warning(msg)
                            warnings.append(msg)
                            writer.writerow([fake_code, "MISSING_COLLEGE", "", "", "", year, round_num, "CET"])
                            placeholders += 1

                    last_code_num = current_code_num
                    logging.info(f"🏫 Found College: {current_college_code} {current_college_name}")

            elif kind == "tbl":
                table = content
                tables_processed += 1

                if current_college_code is None:
                    msg = f"Table before any college header in {os.path.basename(filepath)} (table #{tables_processed}). Dumping."
                    logging.warning(msg)
                    warnings.append(msg)
                    dump_debug(table, f"{Path(filepath).stem}_tbl{tables_processed}_no_header", "no_header")
                    debug_dumps += 1
                    continue

                if len(table.rows) <= 1:
                    continue

                categories = [clean_cell_text(c.text) for c in table.rows[0].cells[1:] if clean_cell_text(c.text)]
                if not categories:
                    msg = f"No categories found for {current_college_code} ({current_college_name}); dumping table."
                    logging.warning(msg)
                    warnings.append(msg)
                    dump_debug(table, current_college_code, "no_categories")
                    debug_dumps += 1
                    continue

                found_any = False
                for row in table.rows[1:]:
                    cells = [clean_cell_text(c.text) for c in row.cells]
                    if not cells or not cells[0]:
                        continue
                    branch = " ".join(cells[0].splitlines()).strip()
                    cutoffs = cells[1:len(categories)+1]

                    for cat, cutoff in zip(categories, cutoffs):
                        if cutoff and cutoff != "--":
                            key = (current_college_code, branch, cat, cutoff)
                            if key in seen_rows:
                                continue
                            writer.writerow([current_college_code, current_college_name, cat, branch, cutoff, year, round_num, "CET"])
                            seen_rows.add(key)
                            rows_written += 1
                            found_any = True

                if not found_any:
                    msg = f"No new rows written for {current_college_code} ({current_college_name}) from table #{tables_processed}."
                    logging.info(msg)

    return {
        "rows": rows_written,
        "tables": tables_processed,
        "dumps": debug_dumps,
        "placeholders": placeholders,
        "colleges": colleges_found,
        "warnings": warnings
    }

# ------------------ Main Script ------------------

def main():
    Path(DEBUG_DIR).mkdir(parents=True, exist_ok=True)

    with open(OUTPUT_FILE, "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(["College_Code", "College_Name", "Category", "Branch", "Cutoff_Rank", "Year", "Round", "Exam_Type"])

    files = sorted([fn for fn in os.listdir(INPUT_DIR) if fn.endswith(".docx") and fn.startswith(("kcetg-round", "kceth-round"))])
    if not files:
        logging.error("No DOCX files found matching expected patterns in INPUT_DIR.")
        return

    total_rows = total_tables = total_dumps = total_placeholders = total_colleges = 0
    all_warnings = []

    for fname in tqdm(files, desc="Processing DOCX files", total=len(files)):
        match = re.match(r"(kcetg|kceth)-round(\d+)-(\d{4})\.docx", fname)
        if not match:
            logging.warning(f"Skipping file with unexpected name format: {fname}")
            continue
        prefix, round_num, year = match.groups()  # ✅ fixed unpacking
        filepath = os.path.join(INPUT_DIR, fname)

        stats = process_docx(filepath, year, round_num)
        total_rows += stats["rows"]
        total_tables += stats["tables"]
        total_dumps += stats["dumps"]
        total_placeholders += stats["placeholders"]
        total_colleges += stats["colleges"]
        all_warnings.extend(stats["warnings"])

        logging.info(f"✅ Wrote {stats['rows']} rows from {fname} (tables={stats['tables']}, colleges={stats['colleges']}, dumps={stats['dumps']})")

    print("\n📊 Summary Report")
    print(f"   Files processed      : {len(files)}")
    print(f"   Colleges detected    : {total_colleges}")
    print(f"   Tables scanned       : {total_tables}")
    print(f"   Rows extracted       : {total_rows}")
    print(f"   Placeholder rows     : {total_placeholders}")
    print(f"   Debug tables dumped  : {total_dumps}")
    print(f"   Warnings logged      : {len(all_warnings)} (see top 10 below)")

    if all_warnings:
        print("\nTop warnings (up to 10):")
        for w in all_warnings[:10]:
            print("  -", w)

    if total_dumps > 0:
        print(f"\nDebug tables saved in: {os.path.abspath(DEBUG_DIR)}")

    logging.info("\n✨ All files processed.")

if __name__ == "__main__":
    main()